#!/usr/bin/env python3
"""Build openegiz/presentation.pptx for the OpenEgiz hackathon."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "presentation.pptx"
IMG = Path("/tmp/pres/img")

# Theme (matches pit UI)
BG = RGBColor(0x0B, 0x0F, 0x14)
PANEL = RGBColor(0x11, 0x16, 0x1D)
LINE = RGBColor(0x22, 0x2A, 0x35)
TEXT = RGBColor(0xE6, 0xED, 0xF3)
MUTED = RGBColor(0x8B, 0x94, 0x9E)
BLUE = RGBColor(0x4E, 0xA1, 0xFF)
GREEN = RGBColor(0x3C, 0xCF, 0x7A)
AMBER = RGBColor(0xFF, 0xB0, 0x20)
ORE = RGBColor(0xC4, 0x7A, 0x35)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W, H = Inches(13.333), Inches(7.5)  # 16:9
MARGIN = Inches(0.55)


def _set_run(run, size=18, bold=False, color=TEXT, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _fill(shape, color: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def _rect(slide, l, t, w, h, color=PANEL):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    _fill(sh, color)
    return sh


def _round(slide, l, t, w, h, color=PANEL):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    _fill(sh, color)
    # Soften corners a bit
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    return sh


def _textbox(slide, l, t, w, h, text, size=18, bold=False, color=TEXT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set(qn("a:anchor"), {
            MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"
        }.get(anchor, "t"))
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def _para(tf, text, size=16, bold=False, color=TEXT, space_before=6, space_after=2,
          align=PP_ALIGN.LEFT):
    p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=bold, color=color)
    return p


def _bullets(slide, l, t, w, h, items, size=17, color=TEXT, bullet_color=BLUE):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(4 if i else 0)
        p.space_after = Pt(6)
        # accent dash
        r0 = p.add_run()
        r0.text = "▸  "
        _set_run(r0, size=size, bold=True, color=bullet_color)
        r1 = p.add_run()
        r1.text = item
        _set_run(r1, size=size, color=color)
    return box


def _header(slide, eyebrow: str, title: str, subtitle: str | None = None):
    _rect(slide, 0, 0, W, Inches(0.08), BLUE)
    _textbox(slide, MARGIN, Inches(0.28), Inches(12), Inches(0.32),
             eyebrow.upper(), size=12, bold=True, color=BLUE)
    _textbox(slide, MARGIN, Inches(0.55), Inches(12.2), Inches(0.55),
             title, size=28, bold=True, color=WHITE)
    if subtitle:
        _textbox(slide, MARGIN, Inches(1.1), Inches(12.2), Inches(0.4),
                 subtitle, size=15, color=MUTED)


def _footer(slide, n: int, total: int):
    _textbox(slide, MARGIN, Inches(7.15), Inches(8), Inches(0.25),
             "Карьер NETA · OpenEgiz v1.0.3 · MIT", size=11, color=MUTED)
    _textbox(slide, Inches(11.2), Inches(7.15), Inches(1.6), Inches(0.25),
             f"{n} / {total}", size=11, color=MUTED, align=PP_ALIGN.RIGHT)


def _kpi(slide, l, t, w, h, label, value, note=None, accent=GREEN):
    card = _round(slide, l, t, w, h, PANEL)
    _rect(slide, l, t, Inches(0.08), h, accent)
    _textbox(slide, l + Inches(0.22), t + Inches(0.18), w - Inches(0.3), Inches(0.28),
             label, size=12, color=MUTED)
    _textbox(slide, l + Inches(0.22), t + Inches(0.48), w - Inches(0.3), Inches(0.5),
             value, size=26, bold=True, color=WHITE)
    if note:
        _textbox(slide, l + Inches(0.22), t + Inches(1.05), w - Inches(0.3), Inches(0.3),
                 note, size=12, color=MUTED)
    return card


def _table(slide, l, t, w, rows, col_widths, header=True):
    """rows: list of list of str. First row is header if header=True."""
    n_rows, n_cols = len(rows), len(rows[0])
    table_shape = slide.shapes.add_table(n_rows, n_cols, l, t, w, Inches(0.38 * n_rows))
    table = table_shape.table
    for i, cw in enumerate(col_widths):
        table.columns[i].width = cw
    for r, row in enumerate(rows):
        for c, cell_text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            run = p.add_run()
            run.text = cell_text
            is_head = header and r == 0
            _set_run(run, size=12 if is_head else 13,
                     bold=is_head or c == 0,
                     color=MUTED if is_head else TEXT)
            # cell fill
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            solid = tcPr.makeelement(qn("a:solidFill"), {})
            srgb = solid.makeelement(qn("a:srgbClr"), {"val": "11161D" if r % 2 == 0 else "161C25"})
            if is_head:
                srgb.set("val", "1A2330")
            solid.append(srgb)
            # clear previous fills
            for child in list(tcPr):
                if child.tag == qn("a:solidFill"):
                    tcPr.remove(child)
            tcPr.insert(0, solid)
    return table_shape


def _picture(slide, path: Path, l, t, w, h=None):
    if not path.exists():
        return None
    if h is None:
        return slide.shapes.add_picture(str(path), l, t, width=w)
    return slide.shapes.add_picture(str(path), l, t, width=w, height=h)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    slides = []

    def add():
        s = prs.slides.add_slide(blank)
        _bg(s)
        slides.append(s)
        return s

    # ------------------------------------------------------------------ 1 TITLE
    s = add()
    _rect(s, 0, 0, W, H, BG)
    hero = IMG / "3d-hero-a.jpg"
    if hero.exists():
        # dimmed hero on the right
        pic = _picture(s, hero, Inches(5.8), Inches(0), Inches(7.6), Inches(7.5))
        if pic:
            # dark overlay panel on left
            pass
    _rect(s, 0, 0, Inches(7.2), H, BG)
    _rect(s, Inches(6.9), 0, Inches(0.35), H, RGBColor(0x0B, 0x0F, 0x14))
    _textbox(s, MARGIN, Inches(1.6), Inches(6.2), Inches(0.35),
             "ХАКАТОН «ЦИФРОВОЙ РУДНИК НА OPENEGIZ»", size=13, bold=True, color=BLUE)
    _textbox(s, MARGIN, Inches(2.1), Inches(6.3), Inches(1.4),
             "Карьер NETA", size=48, bold=True, color=WHITE)
    _textbox(s, MARGIN, Inches(3.5), Inches(6.2), Inches(0.9),
             "Живой цифровой двойник медного карьера:\nдиспетчер · негабарит · допуск после взрыва",
             size=18, color=TEXT)
    _round(s, MARGIN, Inches(4.8), Inches(2.4), Inches(0.55), BLUE)
    _textbox(s, MARGIN, Inches(4.9), Inches(2.4), Inches(0.4),
             "OpenEgiz v1.0.3", size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    _textbox(s, Inches(3.15), Inches(4.9), Inches(3.2), Inches(0.4),
             "MIT · make pit · порт 8095", size=14, color=MUTED)
    _textbox(s, MARGIN, Inches(6.6), Inches(6), Inches(0.3),
             "Команда NETA", size=13, color=MUTED)

    # ------------------------------------------------------------------ 2 GOAL
    s = add()
    _header(s, "01 · задача", "Что решаем",
            "Один связный карьер, три глубоких модуля с замером эффекта на стенде")
    cards = [
        ("Проблема", "Очереди самосвалов, зависания дробилки,\nдолгий и небезопасный допуск после взрыва", AMBER),
        ("Решение", "Цифровой двойник на OpenEgiz:\nсимулятор → MQTT → Ditto → UI / Grafana", BLUE),
        ("Доказательство", "Режим А (модули выкл.) vs Б (вкл.):\n+12 % тоннажа за смену, 0 небезопасных входов", GREEN),
    ]
    for i, (title, body, accent) in enumerate(cards):
        x = MARGIN + i * Inches(4.15)
        _round(s, x, Inches(1.85), Inches(3.95), Inches(3.6), PANEL)
        _rect(s, x, Inches(1.85), Inches(3.95), Inches(0.08), accent)
        _textbox(s, x + Inches(0.3), Inches(2.2), Inches(3.4), Inches(0.45),
                 title, size=20, bold=True, color=accent)
        _textbox(s, x + Inches(0.3), Inches(2.85), Inches(3.4), Inches(2.2),
                 body, size=16, color=TEXT)
    _footer(s, 2, 12)

    # ------------------------------------------------------------------ 3 MINE
    s = add()
    _header(s, "02 · модель", "Карьер как система",
            "Параметры из паспортов техники и открытых источников")
    facts = [
        ("10", "самосвалов Cat 793F", "221 т · 4 354 л"),
        ("2", "экскаватора EX5600-6", "вскрыша + руда"),
        ("165 м", "глубина карьера", "этапы −195…−255 м"),
        ("12 ч", "смена", "взрыв в 06:00"),
        ("0,35 %", "содержание Cu", "извлечение 88 %"),
        ("36", "двойников в Ditto", "neta.pit:*"),
    ]
    for i, (v, t1, t2) in enumerate(facts):
        col, row = i % 3, i // 3
        x = MARGIN + col * Inches(4.15)
        y = Inches(1.85) + row * Inches(2.2)
        _round(s, x, y, Inches(3.95), Inches(2.0), PANEL)
        _textbox(s, x + Inches(0.3), y + Inches(0.3), Inches(3.4), Inches(0.55),
                 v, size=32, bold=True, color=BLUE)
        _textbox(s, x + Inches(0.3), y + Inches(0.95), Inches(3.4), Inches(0.35),
                 t1, size=16, bold=True, color=WHITE)
        _textbox(s, x + Inches(0.3), y + Inches(1.35), Inches(3.4), Inches(0.35),
                 t2, size=13, color=MUTED)
    _footer(s, 3, 12)

    # ------------------------------------------------------------------ 4 ARCH
    s = add()
    _header(s, "03 · архитектура", "Данные идут через платформу",
            "Не отдельное приложение «для галочки», а поток MQTT → Ditto → UI")
    steps = [
        ("1", "Симулятор\npit-live", "Discrete-event\nсмена × speedup"),
        ("2", "MQTT\nMosquitto", "Телеметрия\nи команды"),
        ("3", "Eclipse\nDitto", "36 двойников\nneta.pit:*"),
        ("4", "UI 8095\n+ Grafana", "SSE / InfluxDB\nкарта и KPI"),
    ]
    for i, (n, title, body) in enumerate(steps):
        x = MARGIN + i * Inches(3.1)
        _round(s, x, Inches(2.0), Inches(2.85), Inches(3.4), PANEL)
        _round(s, x + Inches(1.05), Inches(2.25), Inches(0.7), Inches(0.7), BLUE)
        _textbox(s, x + Inches(1.05), Inches(2.35), Inches(0.7), Inches(0.5),
                 n, size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        _textbox(s, x + Inches(0.2), Inches(3.2), Inches(2.45), Inches(0.9),
                 title, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        _textbox(s, x + Inches(0.2), Inches(4.2), Inches(2.45), Inches(0.9),
                 body, size=13, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 3:
            _textbox(s, x + Inches(2.7), Inches(3.35), Inches(0.4), Inches(0.4),
                     "→", size=22, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    _textbox(s, MARGIN, Inches(5.7), Inches(12.2), Inches(0.8),
             "Сервисы поверх Ditto: dispatcher-01 · crusher-guard-01 · blast-control-01 · advisor-01\n"
             "Режимы А/Б и сценарии «что если» — свойства двойников control-01 / whatif-01.",
             size=14, color=MUTED)
    _footer(s, 4, 12)

    # ------------------------------------------------------------------ 5 MODULES
    s = add()
    _header(s, "04 · модули", "Три модуля, один карьер",
            "Каждый модуль: проблема → решение → метрика → замер на стенде")
    mods = [
        ("М1", "Диспетчеризация", "LP-сервис vs фиксированный\nраздел 5/5",
         "+1 642 т/смену\n(+3,0 %)", "очередь 11,8 → 9,5 ч", BLUE),
        ("М2", "Негабарит", "Камера + гидромолот\nдо зависания дробилки",
         "простой 23 → 1 мин", "−22 мин простоя", AMBER),
        ("М3", "Допуск после взрыва", "Дрон 24 ячейки vs\nручной замер ВГСЧ",
         "+4 310 т/смену\n(+7,8 %)", "0 небезопасных входов", GREEN),
    ]
    for i, (code, title, how, effect, metric, accent) in enumerate(mods):
        x = MARGIN + i * Inches(4.15)
        _round(s, x, Inches(1.85), Inches(3.95), Inches(4.6), PANEL)
        _rect(s, x, Inches(1.85), Inches(3.95), Inches(0.1), accent)
        _textbox(s, x + Inches(0.3), Inches(2.15), Inches(3.4), Inches(0.35),
                 code, size=14, bold=True, color=accent)
        _textbox(s, x + Inches(0.3), Inches(2.5), Inches(3.4), Inches(0.45),
                 title, size=20, bold=True, color=WHITE)
        _textbox(s, x + Inches(0.3), Inches(3.15), Inches(3.4), Inches(0.9),
                 how, size=14, color=MUTED)
        _textbox(s, x + Inches(0.3), Inches(4.2), Inches(3.4), Inches(0.3),
                 "ЭФФЕКТ vs А", size=11, bold=True, color=MUTED)
        _textbox(s, x + Inches(0.3), Inches(4.5), Inches(3.4), Inches(0.85),
                 effect, size=20, bold=True, color=accent)
        _textbox(s, x + Inches(0.3), Inches(5.55), Inches(3.4), Inches(0.5),
                 metric, size=14, color=TEXT)
    _footer(s, 5, 12)

    # ------------------------------------------------------------------ 6 RESULTS
    s = add()
    _header(s, "05 · замер", "А vs Б: 30 смен × 12 ч",
            "Одинаковые seeds 1…30 — сравнение в равных условиях")
    _kpi(s, MARGIN, Inches(1.75), Inches(4.0), Inches(1.45),
         "Тоннаж за смену", "+6 650 т", "+12,1 % к режиму А", GREEN)
    _kpi(s, MARGIN + Inches(4.2), Inches(1.75), Inches(4.0), Inches(1.45),
         "Медь в руде", "+12,5 т Cu", "≈ +119 тыс. USD / смену*", BLUE)
    _kpi(s, MARGIN + Inches(8.4), Inches(1.75), Inches(4.0), Inches(1.45),
         "Безопасность", "0 входов в газ", "доля небезопасных 0,20 → 0", AMBER)

    rows = [
        ["Конфиг", "т/смену", "Руда, т", "Простой др., мин", "Забой закр., мин"],
        ["А · все выкл.", "55 025 ± 2 592", "19 508", "23", "153"],
        ["М1 · диспетчер", "56 667 (+3 %)", "21 746", "30", "153"],
        ["М2 · камера", "55 366 (+0,6 %)", "20 086", "1", "153"],
        ["М3 · дрон", "59 335 (+7,8 %)", "20 598", "23", "131"],
        ["Б · все вкл.", "61 676 (+12 %)", "23 582", "6", "131"],
    ]
    _table(s, MARGIN, Inches(3.5), Inches(12.2), rows,
           [Inches(2.6), Inches(2.6), Inches(2.2), Inches(2.5), Inches(2.3)])
    _textbox(s, MARGIN, Inches(6.55), Inches(12.2), Inches(0.35),
             "* 9 500 USD/т Cu (LME 2025, диапазон 9–10 тыс.). Команда: make pit-experiment SHIFTS=30",
             size=11, color=MUTED)
    _footer(s, 6, 12)

    # ------------------------------------------------------------------ 7 VIS 2D
    s = add()
    _header(s, "06 · визуализация", "2D-карта карьера в реальном времени",
            "Движение техники из двойников Ditto (SSE), не из локальной базы")
    pic = IMG / "h-2d.png"
    if pic.exists():
        # crop-friendly fit under header
        _picture(s, pic, Inches(0.4), Inches(1.55), Inches(12.5), Inches(5.35))
    else:
        _textbox(s, MARGIN, Inches(3), Inches(12), Inches(1),
                 "Скриншот недоступен — откройте http://127.0.0.1:8095/", size=16, color=MUTED)
    _footer(s, 7, 12)

    # ------------------------------------------------------------------ 8 VIS 3D
    s = add()
    _header(s, "06 · визуализация", "3D-вид: забои, рампы, отвалы, дробилка",
            "Тот же поток двойников, что и на 2D-карте")
    left = IMG / "3d-hero-a.jpg"
    right = IMG / "3d-shovel-a.jpg"
    if left.exists():
        _picture(s, left, Inches(0.4), Inches(1.65), Inches(7.6), Inches(4.9))
    if right.exists():
        _picture(s, right, Inches(8.2), Inches(1.65), Inches(4.7), Inches(4.9))
    _footer(s, 8, 12)

    # ------------------------------------------------------------------ 9 WHATIF
    s = add()
    _header(s, "07 · сценарии", "«Что если» на живом стенде",
            "Тумблеры whatif-01 меняют условия сразу; А/Б сравниваются на тех же seeds")
    rows = [
        ["Сценарий", "А, т/смену", "Б, т/смену", "Б − А"],
        ["База", "57 054", "63 577", "+6 523"],
        ["Встал shovel-01", "20 569 (−64 %)", "34 023", "+13 453"],
        ["3 самосвала в ремонте", "37 806 (−34 %)", "51 558", "+13 752"],
        ["Плохой взрыв (x50 0,30 м)", "52 213 (−8 %)", "60 054", "+7 841"],
        ["Штиль при взрыве", "44 150 (−23 %)", "42 048", "−2 102*"],
        ["Карьер +6 лет (−255 м)", "51 164 (−10 %)", "58 002", "+6 838"],
    ]
    _table(s, MARGIN, Inches(1.75), Inches(12.2), rows,
           [Inches(3.8), Inches(2.8), Inches(2.8), Inches(2.8)])
    _textbox(s, MARGIN, Inches(5.9), Inches(12.2), Inches(0.9),
             "* При штиле дрон ждёт ниже ПДК и не пускает раньше закона (≥ 30 мин) — "
             "тоннаж чуть ниже, зато 0 мин в газе выше ПДК (режим А: 20,7 мин).\n"
             "Также: дождь, закрытие дороги на отвал, остановка дробилки/фабрики, бульдозер.",
             size=13, color=MUTED)
    _footer(s, 9, 12)

    # ------------------------------------------------------------------ 10 CALIB
    s = add()
    _header(s, "08 · достоверность", "Калибровка по паспортам и полевым данным",
            "openpit.validate — 10 смен, режим Б")
    rows = [
        ["Метрика", "Модель", "Реальный диапазон", "Вердикт"],
        ["Скорость на рампе 10 %", "12,9 км/ч", "12,0 – 12,9", "на границе"],
        ["Число ковшей (221 т)", "5", "4 – 5", "в диапазоне"],
        ["Загрузка", "2,75 мин", "2,5 – 3,1", "в диапазоне"],
        ["Разгрузка", "1,10 мин", "1,0 – 2,25", "в диапазоне"],
        ["Расход дизеля", "149 л/ч", "121 – 240", "в диапазоне"],
        ["Удельный расход", "0,28 л/т", "0,25 – 0,42", "в диапазоне"],
        ["Эффективное время смены", "8,3 ч", "6,6 – 9,0", "в диапазоне"],
    ]
    _table(s, MARGIN, Inches(1.75), Inches(12.2), rows,
           [Inches(3.8), Inches(2.4), Inches(3.2), Inches(2.8)])
    _textbox(s, MARGIN, Inches(5.85), Inches(12.2), Inches(0.9),
             "Источники: Cat 793F Specalog / Performance Handbook; Hitachi EX5600-6; "
             "ПБ 13-407-01; GN 2.2.5.3532-18 (ПДК CO/NO₂);\n"
             "KAZ Minerals (Aktogay/Bozshakol) — содержание и извлечение; LME — цена меди.",
             size=13, color=MUTED)
    _footer(s, 10, 12)

    # ------------------------------------------------------------------ 11 RUN
    s = add()
    _header(s, "09 · запуск", "Судья поднимает стенд по README",
            "Чистая Ubuntu 24.04 · Docker · без ручных шагов")
    _round(s, MARGIN, Inches(1.85), Inches(7.8), Inches(4.5), PANEL)
    _textbox(s, MARGIN + Inches(0.4), Inches(2.15), Inches(7.0), Inches(0.4),
             "Команды", size=16, bold=True, color=BLUE)
    cmds = [
        "git clone <repo> && cd openegiz",
        "cp deploy/compose/.env.example deploy/compose/.env",
        "make pit          # платформа + карьер (~1–2 мин)",
        "open http://127.0.0.1:8095/",
        "make pit-experiment SHIFTS=30   # таблица А/Б",
    ]
    box = s.shapes.add_textbox(MARGIN + Inches(0.4), Inches(2.7), Inches(7.0), Inches(3.2))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(cmds):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = line
        _set_run(run, size=16, color=GREEN, font="Consolas")

    side = [
        ("UI", "http://127.0.0.1:8095/"),
        ("Ditto", "http://127.0.0.1:8081/"),
        ("Grafana", "http://127.0.0.1:3000/"),
        ("Лицензия", "MIT"),
    ]
    for i, (k, v) in enumerate(side):
        y = Inches(1.85) + i * Inches(1.1)
        _round(s, Inches(8.7), y, Inches(4.0), Inches(0.95), PANEL)
        _textbox(s, Inches(8.95), y + Inches(0.15), Inches(3.5), Inches(0.3),
                 k, size=12, color=MUTED)
        _textbox(s, Inches(8.95), y + Inches(0.42), Inches(3.5), Inches(0.35),
                 v, size=15, bold=True, color=WHITE)
    _footer(s, 11, 12)

    # ------------------------------------------------------------------ 12 END
    s = add()
    _rect(s, 0, 0, W, H, BG)
    _rect(s, 0, 0, W, Inches(0.08), BLUE)
    _textbox(s, MARGIN, Inches(1.8), Inches(12.2), Inches(0.4),
             "ИТОГ", size=14, bold=True, color=BLUE)
    _textbox(s, MARGIN, Inches(2.3), Inches(12.2), Inches(1.0),
             "Один карьер. Три модуля. Доказанный эффект.",
             size=32, bold=True, color=WHITE)
    points = [
        "+12 % тоннажа за смену при включённых модулях (Б vs А, n=30)",
        "0 небезопасных входов в газ; простой дробилки −17 мин",
        "36 двойников Ditto, живая 2D/3D-карта, Grafana, сценарии «что если»",
        "Запуск: make pit → http://127.0.0.1:8095/",
    ]
    _bullets(s, MARGIN, Inches(3.6), Inches(12), Inches(2.5), points, size=18, bullet_color=GREEN)
    _textbox(s, MARGIN, Inches(6.4), Inches(12.2), Inches(0.4),
             "Спасибо · команда NETA · OpenEgiz v1.0.3", size=14, color=MUTED)

    # footers for title? skip. Fix total count on earlier slides if needed — we used 12.
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"wrote {path} ({path.stat().st_size // 1024} KB)")
